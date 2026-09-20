"""Offline inspection of reviewed raw receipts. No claims of engine eligibility.

Writes a new immutable profile directory. CSV and XML parsers are stdlib;
read-only XLSX inspection requires openpyxl. No workbook is modified, recalculated,
or executed. Formula strings remain strings. No imputation or inferred zeroes.
"""
from collections import Counter
import csv
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
import re
from pathlib import Path
import argparse
from uuid import uuid4
import xml.etree.ElementTree as ET
from zipfile import ZipFile

PARSER_VERSION = "research-profile-v1.1"


def reported_cell(text):
    """Keep ranges/censoring as reported, never substitute a midpoint or zero."""
    value = text.strip()
    base = {"original_text": text, "unit": "wt_pct"}
    if value in ("", "-"):
        return {**base, "kind": "UNSPECIFIED", "value": None, "missing_reason": "Source blank/dash; not assumed zero"}
    if re.fullmatch(r"\d+(?:\.\d+)?–\d+(?:\.\d+)?", value):
        lower, upper = value.split("–")
        return {**base, "kind": "REPORTED_RANGE", "lower": lower, "upper": upper, "value": None}
    if re.fullmatch(r"<\d+(?:\.\d+)?", value):
        return {**base, "kind": "REPORTED_UPPER_BOUND", "upper_exclusive": value[1:], "value": None}
    try:
        number = Decimal(value)
        if not number.is_finite():
            raise ValueError("NONFINITE_TABLE_VALUE")
        return {**base, "kind": "REPORTED_VALUE", "value": str(number)}
    except InvalidOperation:
        raise ValueError("UNREVIEWED_TABLE_VALUE") from None


def fabris_entities(tables):
    material_table = next(t for t in tables if t["label"] == "Table 2")
    recipe_table = next(t for t in tables if t["label"] == "Table 1")
    common = {"evidence_kind": "OBSERVED", "qualifier": "REPORTED", "core_engine_eligible": False,
              "status": "QUARANTINED_FOR_CHEMISTRY", "analysis_basis": "UNKNOWN",
              "limitation": "Product/lot and analysis basis unresolved; ranges/dashes are not exact chemistry."}
    materials = [{**common, "source_record_id": "fabris-2024:table2:" + row[0], "name_as_reported": row[0],
                  "composition": {header: reported_cell(text) for header, text in zip(material_table["rows"][0][1:], row[1:])}}
                 for row in material_table["rows"][1:]]
    recipes = [{**common, "source_record_id": "fabris-2024:table1:" + name, "name_as_reported": name,
                "ingredients": [{"material_as_reported": row[0], "amount": reported_cell(row[column])}
                                for row in recipe_table["rows"][1:]],
                "limitation": "Published ingredient ranges, not exact reproducible recipe; no midpoint substitution."}
               for column, name in enumerate(recipe_table["rows"][0][1:], 1)]
    return materials, recipes


def matrix_profile(rows):
    nonempty = [row for row in rows if any(v is not None and str(v).strip() for v in row)]
    counts = Counter(json.dumps(row, ensure_ascii=False, default=str) for row in nonempty)
    return {"nonempty_rows_including_headers_notes": len(nonempty),
            "max_columns": max(map(len, rows), default=0),
            "exact_duplicate_nonempty_rows": sum(n - 1 for n in counts.values()),
            "blank_cells": sum(v is None or (isinstance(v, str) and not v.strip()) for row in rows for v in row),
            "interpretation": "Matrix rows include headers/notes; NOT a count of independent experiments."}


def profile_csv(raw):
    try:
        content = raw.decode("utf-8-sig")
        encoding = "utf-8-sig"
    except UnicodeDecodeError:
        content = raw.decode("cp1252")
        encoding = "cp1252 (explicit fallback)"
    # Candidate delimiter comes from the observed header, never decimal replacement.
    header = content.splitlines()[0]
    delimiter = ";" if header.count(";") > header.count(",") else ","
    rows = list(csv.reader(io.StringIO(content), delimiter=delimiter))
    return {"encoding": encoding, "delimiter": delimiter, "profile": matrix_profile(rows), "rows": rows}


def uci_records(raw):
    table = profile_csv(raw)
    rows = table["rows"]
    expected = ["Ceramic Name", "Part", "Na2O", "MgO", "Al2O3", "SiO2", "K2O", "CaO", "TiO2", "Fe2O3",
                "MnO", "CuO", "ZnO", "PbO2", "Rb2O", "SrO", "Y2O3", "ZrO2", "P2O5"]
    if rows[0] != expected:
        raise ValueError("UCI_SCHEMA_CHANGED")
    records = []
    for index, row in enumerate(rows[1:], 2):
        if len(row) != len(expected) or row[1] not in ("Body", "Glaze"):
            raise ValueError("UCI_ROW_SHAPE_OR_PART_CHANGED")
        measurements = []
        for column, text in zip(expected[2:], row[2:]):
            unit = "wt_pct" if column in expected[2:10] else "ppm"
            try:
                number = Decimal(text)
                if not number.is_finite():
                    raise ValueError("UCI_INVALID_MEASUREMENT")
            except InvalidOperation:
                raise ValueError("UCI_NONNUMERIC_VALUE") from None
            measurements.append({"reported_as": column, "reported_value": text, "reported_unit": unit,
                                 "wt_pct": str(number if unit == "wt_pct" else number / 10000) if number >= 0 else None,
                                 "conversion": "identity" if unit == "wt_pct" else "ppm / 10000",
                                 "quality_flag": "NEGATIVE_REPORTED_VALUE_REVIEW" if number < 0 else None,
                                 "oxide_identity_review": "REQUIRED" if column == "PbO2" else "AS_REPORTED"})
        records.append({"source_record_id": f"uci-583:csv-row-{index}", "source_row": index,
                        "ceramic_name": row[0], "part": row[1], "measurements": measurements,
                        "evidence_kind": "OBSERVED", "method_kind": "EMPIRICAL", "qualifier": "REPORTED",
                        "measurement_method": "EDXRF", "core_engine_eligible": False,
                        "limitation": "Fired archaeological specimen; not an unfired material or a recipe.",
                        "independent_specimen_id": None, "uncertainty": None})
    return table, records


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--storage", type=Path, required=True)
    args = parser.parse_args()
    root = args.storage.resolve()
    # Use completed sources only; deduplicate identical content across receipts.
    items = {}
    source_meta = {}
    for receipt_path in sorted((root / "receipts").glob("*.json"), key=lambda p: json.loads(p.read_text(encoding="utf-8"))["created_at"]):
        receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
        source_meta.update(receipt["sources"])
        for item in receipt["files"]:
            if item["source_id"] in receipt["sources"]:
                items[(item["source_id"], item["filename"], item["sha256"])] = item
    outputs, summary = [], []
    for item in items.values():
        path = (root / item["raw_path"]).resolve()
        if not path.is_relative_to(root):
            raise ValueError("UNSAFE_RAW_PATH")
        raw = path.read_bytes()
        if hashlib.sha256(raw).hexdigest() != item["sha256"]:
            raise ValueError("RAW_CHECKSUM_MISMATCH")
        source, name = item["source_id"], item["filename"]
        payload = None
        if source == "uci-583" and name.endswith(".zip"):
            with ZipFile(io.BytesIO(raw)) as archive:
                members = [m for m in archive.infolist() if m.filename.endswith(".csv")]
                if len(members) != 1 or members[0].file_size > 1_000_000:
                    raise ValueError("UCI_ARCHIVE_SHAPE_CHANGED")
                table, records = uci_records(archive.read(members[0]))
                payload = {"table": table, "records": records}
                summary.append({"source_id": source, "kind": "FIRED_SAMPLE_MEASUREMENT_ROWS", "count": len(records),
                                "parts": dict(Counter(r["part"] for r in records)),
                                "exact_duplicate_data_rows": len(records) - len({tuple(r) for r in table["rows"][1:]}),
                                "negative_cells": sum(m["quality_flag"] is not None for r in records for m in r["measurements"]),
                                "rows_with_negative_cells": sum(any(m["quality_flag"] for m in r["measurements"]) for r in records)})
        elif name.endswith(".csv"):
            payload = profile_csv(raw)
            summary.append({"source_id": source, "filename": name, **payload["profile"]})
        elif name.endswith(".xlsx"):
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            sheets = []
            for sheet in workbook:
                rows = [[v.isoformat() if hasattr(v, "isoformat") else v for v in row] for row in sheet.iter_rows(values_only=True)]
                sheets.append({"name": sheet.title, "rows": rows, "profile": matrix_profile(rows)})
            workbook.close()
            payload = {"sheets": sheets, "note": "Cell matrices only; multirow headers and notes are not experiments."}
            summary.append({"source_id": source, "filename": name, "sheets": [{"name": s["name"], **s["profile"]} for s in sheets]})
        elif source == "fabris-2024" and name.endswith(".xml"):
            doc = ET.fromstring(raw)
            tables = []
            for table in doc.findall(".//table-wrap"):
                tables.append({"id": table.get("id"), "label": "".join(table.find("label").itertext()),
                               "caption": "".join(table.find("caption").itertext()),
                               "rows": [["".join(cell.itertext()) for cell in row] for row in table.findall(".//tr")],
                               "cell_attributes": [[dict(cell.attrib) for cell in row] for row in table.findall(".//tr")],
                               "grid_note": "Rows follow XML elements; rowspan/colspan retained, not expanded. Do not infer aligned columns from ragged rows."})
            materials, recipes = fabris_entities(tables)
            payload = {"tables": tables, "research_materials": materials, "research_recipe_ranges": recipes,
                       "notes": "Preserved table text. Dashes remain dashes; anonymous products are not identified brands."}
            summary.append({"source_id": source, "kind": "PUBLISHED_TABLES", "count": len(tables),
                            "research_material_rows": len(materials), "recipe_range_sets": len(recipes),
                            "core_engine_eligible": 0})
        elif source == "kiln-controller" and name.endswith(".json"):
            payload = {"reported_profile": json.loads(raw), "evidence_kind": "OBSERVED", "qualifier": "REPORTED",
                       "record_kind": "SOFTWARE_EXAMPLE_NOT_MEASURED_FIRING", "temperature_unit": None, "time_unit": None,
                       "units_status": "Requires per-profile confirmation; repository display defaults are not per-file metadata.",
                       "safe_to_execute": False, "core_engine_eligible": False}
            summary.append({"source_id": source, "filename": name, "kind": "UNVALIDATED_PROGRAM", "points": len(payload["reported_profile"]["data"])})
        if payload is not None:
            outputs.append({"source_id": source, "filename": name, "source_metadata": source_meta[source],
                            "raw_sha256": item["sha256"], "source_url": item["source_url"],
                            "parser_version": PARSER_VERSION, "payload": payload})
    target = root / "profiles" / uuid4().hex
    target.mkdir(parents=True, exist_ok=False)
    for index, output in enumerate(outputs):
        (target / f"{index:03d}-{output['source_id']}.json").write_text(json.dumps(output, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    report = {"parser_version": PARSER_VERSION, "parser_code_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
              "created_at": datetime.now(timezone.utc).isoformat(),
              "source_count": len(source_meta), "profiled_files": len(outputs), "summary": summary,
              "production_ready_material_analyses": 0, "training_release": False}
    (target / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"profile_directory": str(target), "report": report}, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
